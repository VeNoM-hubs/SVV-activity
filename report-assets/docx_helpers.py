"""Shared python-docx helpers for the Team 04 test reports."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "report-assets"
BUGS = ROOT / "assets"

TEAM = [
    ("Bhaskar Lukram", "RA2311033010003", "Test planning, test scenarios, Boundary Value Analysis"),
    ("Khushi Raghav", "RA2311033010015", "Equivalence Class Partitioning, Cause–Effect Graph, Decision Table"),
    ("Kartik Gupta", "RA2311033010045", "White-box testing, code coverage, unit & integration testing"),
    ("Ayush Pandey", "RA2311033010044", "System & acceptance testing, defect reporting"),
]

DARK = RGBColor(0x1F, 0x2A, 0x37)
ACCENT = RGBColor(0x1A, 0x7F, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1F2A37"
ALT_FILL = "F3F5F7"
PASS_FILL = "DDF4E4"
FAIL_FILL = "FDE2E1"
KEY_FILL = "E9EEF3"

fig_no = 0
tab_no = 0


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, color=None, size=10, align=None, font=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        if font:
            run.font.name = font
        if align:
            p.alignment = align
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def keep_together(t, widths=None):
    """Stop a table from splitting across pages and pin its column widths."""
    t.autofit = False
    last = len(t.rows) - 1
    for r_i, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if r_i < last:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
        if widths:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)


def repeat_header(t):
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def status_fill(cell, text):
    t = text.upper()
    if t.startswith("PASS"):
        shade(cell, PASS_FILL)
    elif t.startswith("FAIL"):
        shade(cell, FAIL_FILL)


def add_caption(doc, kind, text):
    global fig_no, tab_no
    if kind == "Table":
        tab_no += 1
        n = tab_no
    else:
        fig_no += 1
        n = fig_no
    cap = doc.add_paragraph(style="Caption")
    cap.add_run(f"{kind} {n}: {text}")
    return cap


def table(doc, header, rows, widths=None, caption=None, status_col=None, size=9.5,
          together=True, gap=True):
    if caption:
        add_caption(doc, "Table", caption).paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        set_cell(c, h, bold=True, color=WHITE, size=size)
        shade(c, HEADER_FILL)
    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell(cells[i], val, size=size)
            if r_i % 2 == 1:
                shade(cells[i], ALT_FILL)
            if status_col is not None and i == status_col:
                status_fill(cells[i], val)
    if together:
        keep_together(t, widths)
    else:
        t.autofit = False
        repeat_header(t)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for i, w in enumerate(widths or []):
                row.cells[i].width = Inches(w)
        for i, w in enumerate(widths or []):
            t.columns[i].width = Inches(w)
    if gap:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def kv_table(doc, pairs, key_width=1.7, val_width=4.8, status_key="Status", container=None, size=10):
    t = (container or doc).add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in pairs:
        cells = t.add_row().cells
        set_cell(cells[0], k, bold=True, size=size)
        shade(cells[0], KEY_FILL)
        set_cell(cells[1], v, size=size)
        if k == status_key:
            status_fill(cells[1], v)
            for r in cells[1].paragraphs[0].runs:
                r.bold = True
    keep_together(t, [key_width, val_width])
    if container is None:
        doc.add_paragraph()
    return t


def defect_block(doc, pairs, img, caption, left_w=4.35, right_w=2.0, img_w=1.85, size=9):
    """Defect details (left) next to the screenshot evidence (right)."""
    global fig_no
    outer = doc.add_table(rows=1, cols=2)
    outer.autofit = False
    left, right = outer.rows[0].cells
    left.width, right.width = Inches(left_w), Inches(right_w)
    outer.columns[0].width, outer.columns[1].width = Inches(left_w), Inches(right_w)
    outer.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    kv_table(doc, pairs, key_width=1.25, val_width=left_w - 1.45, container=left, size=size)
    left._tc.remove(left.paragraphs[0]._p)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(img), width=Inches(img_w))
    fig_no += 1
    cap = right.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(f"Figure {fig_no}: {caption}")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, path, caption, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = add_caption(doc, "Figure", caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return fig_no


def para(doc, text, bold_lead=None, italic=False, size=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        if size:
            r.font.size = Pt(size)
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return p


def bullets(doc, items, style="List Bullet", size=None):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2)
        parts = [(it[0], True), (it[1], False)] if isinstance(it, tuple) else [(it, False)]
        for text, bold in parts:
            r = p.add_run(text)
            r.bold = bold
            if size:
                r.font.size = Pt(size)


def code_block(doc, lines, size=9.5):
    for ln in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.left_indent = Cm(0.4)
        pf.keep_with_next = True
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F3F5F7")
        p_pr.append(shd)
        r = p.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(size)
    p.paragraph_format.keep_with_next = False
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_field(paragraph, instr):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r._r.append(b)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    paragraph.add_run("Right-click and choose Update Field to build the table of contents.")
    r4 = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r4._r.append(e)


def crop_slide(n):
    src = FIG / f"slide-{n}.png"
    dst = FIG / f"fig-slide-{n}.png"
    with Image.open(src) as im:
        im.crop((56, 48, 1864, 1032)).save(dst)
    return dst


def new_document(body_size=11, h_sizes=(16, 13, 11.5), margins=(0.95, 0.9), space_after=6):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(margins[0])
    sec.top_margin = sec.bottom_margin = Inches(margins[1])

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(body_size)
    normal.paragraph_format.space_after = Pt(space_after)
    for lvl, size in zip((1, 2, 3), h_sizes):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = DARK if lvl == 1 else ACCENT
        st.font.bold = True
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.color.rgb = RGBColor(0x55, 0x5F, 0x6B)

    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run("Team 04 · Manual Test Report · Online Food Ordering System · Page ")
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    pg = OxmlElement("w:fldSimple")
    pg.set(qn("w:instr"), "PAGE")
    footer_p._p.append(pg)
    return doc


def cover(doc, top_gap=3):
    for _ in range(top_gap):
        doc.add_paragraph()
    for text, size, color, bold, italic in (
        ("MANUAL TEST REPORT", 28, DARK, True, False),
        ("Online Food Ordering System", 20, ACCENT, True, False),
        ("Manual Test Case Design & Testing Levels", 13, None, False, False),
        ("Black-box design (BVA, ECP, Cause–Effect Graph, Decision Table) · "
         "White-box testing & source code coverage · Unit, Integration, System & Acceptance testing",
         10, None, False, True),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TEAM 4")
    r.bold = True
    r.font.size = Pt(16)

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(("Name", "Register Number (RA)", "Responsibility")):
        set_cell(t.rows[0].cells[i], h, bold=True, color=WHITE, size=10.5)
        shade(t.rows[0].cells[i], HEADER_FILL)
    for name, ra, role in TEAM:
        c = t.add_row().cells
        set_cell(c[0], name, bold=True, size=10.5)
        set_cell(c[1], ra, size=10.5)
        set_cell(c[2], role, size=9.5)
    keep_together(t, [1.7, 1.9, 2.8])

    doc.add_paragraph()
    kv_table(doc, [
        ("Application", "FOODIE: Online Food Ordering System (test build)"),
        ("Document type", "Manual Test Report (test design, execution, defects)"),
        ("Version / Date", "1.0 · October 2026"),
        ("Total test cases", "10 executed · 7 passed · 3 failed · pass rate 70%"),
        ("Defects", "3 (1 Critical, 2 Major)"),
    ], key_width=1.8, val_width=4.6)


def team_table(doc, size=10):
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(("RA Number", "Student Name", "Role in testing")):
        set_cell(t.rows[0].cells[i], h, bold=True, color=WHITE, size=size)
        shade(t.rows[0].cells[i], HEADER_FILL)
    for name, ra, role in TEAM:
        c = t.add_row().cells
        set_cell(c[0], ra, size=size)
        set_cell(c[1], name, size=size)
        set_cell(c[2], role, size=size - 0.5)
    keep_together(t, [1.7, 1.6, 3.2])
    return t
